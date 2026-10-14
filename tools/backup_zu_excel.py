#!/usr/bin/env python3
"""Einmalige Übernahme einer Sicherung der Wartungs-App in die Excel-Arbeitsmappe „Wartungsdaten.xlsx“.

Erzeugt je Zielblatt eine Tab-getrennte Textdatei nach „PWA_Importformat.md“ (Stand 14.10.2026),
die im Editor kopiert und in Excel unter die jeweilige Tabelle eingefügt wird.

Aufruf (benötigt: pip install openpyxl):
    python backup_zu_excel.py
        nimmt die neueste Wartung_Backup*.json und Wartungsdaten*.xlsx aus dem aktuellen Ordner
    python backup_zu_excel.py Wartung_Backup_2026-10-14.json --excel Wartungsdaten.xlsx

Vorher wird der Aufbau der Mappe geprüft (Tabellen, Spaltennamen, lückenlose Eingabeblöcke). Passt er nicht,
bricht das Programm ohne Ausgabe ab und schreibt einen Bericht nach „Strukturpruefung.txt“.
Danach wird abgeglichen: Termine, Mitarbeiter und Technische Daten, die es schon gibt, werden nicht erneut
ausgegeben; die Kunden-Nr neuer Aggregate stammt aus dem Blatt „Kunden“.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

MAX_ZELLE = 32_767
GEKUERZT = ' … (gekürzt, siehe Protokoll)'

LEISTUNGEN = ['elektr_wartung', 'motor_wartung', 'oelwechsel', 'dguv',
              'kuehlmittel_intern', 'kuehlmittel_extern', 'luftfilter', 'zahnriemenwechsel']
INTERVALLE = ['Elektr_Wartung_Intervall', 'Motorwartung_Intervall', 'Oelwechsel_Intervall', 'DGUV_Intervall',
              'Kuehlmittel_Intervall_intern', 'Kuehlmittel_Intervall_extern', 'Luftfilter_Intervall']


# ── Feldformate (Abschnitt 2 und 3) ─────────────────────────

def text(wert) -> str:
    if wert is None:
        return ''
    t = str(wert).replace('\r\n', ' / ').replace('\n', ' / ').replace('\r', ' / ').replace('\t', ' ').strip()
    # Excel würde sonst eine Formel daraus machen
    return "'" + t if t[:1] in ('=', '+', '-', '@') else t


def zahl(wert) -> float | None:
    if wert is None or wert == '':
        return None
    if isinstance(wert, (int, float)):
        return float(wert)
    try:
        return float(str(wert).strip().replace('.', '').replace(',', '.')
                     if ',' in str(wert) else str(wert).strip())
    except ValueError:
        return None


def ganzzahl(wert) -> str:
    z = zahl(wert)
    return '' if z is None else str(round(z))


def dezimal(wert) -> str:
    z = zahl(wert)
    if z is None:
        return ''
    return (str(int(z)) if z.is_integer() else repr(z)).replace('.', ',')


def geld(wert) -> str:
    z = zahl(wert)
    return '' if z is None else f'{z:.2f}'.replace('.', ',')


def datum(iso) -> str:
    if not iso:
        return ''
    return datetime.strptime(str(iso)[:10], '%Y-%m-%d').strftime('%d.%m.%Y')


def ja_nein(wert) -> str:
    return '' if wert is None else ('Ja' if wert else 'Nein')


def x(wert) -> str:
    return 'x' if wert else ''


def ref(anzeige: str, schluessel: str, leer: str) -> str:
    anzeige = text(anzeige).replace('[', '(').replace(']', ')') or leer
    return f'{anzeige} [{schluessel}]'


def ref_aggregat(bezeichnung, kommission) -> str:
    return ref(bezeichnung, kommission, '(ohne Bezeichnung)')


def ref_kunde(name, nr) -> str:
    return ref(name, nr, '(ohne Namen)')


def schluessel_aus_ref(wert) -> str:
    """Text zwischen dem letzten „[“ und dem abschließenden „]“ (wie die Excel-Formel)."""
    t = str(wert or '').strip()
    if t.endswith(']') and '[' in t:
        return t[t.rindex('[') + 1:-1].strip()
    return t


def server_link(pfad) -> str:
    """Format G:/Überordner\\Unterordner\\Ziel – nach dem Laufwerk ein „/“, danach nur „\\“."""
    t = str(pfad or '').strip().strip('"')
    m = re.match(r'^([A-Za-z]):[\\/]*(.*)$', t)
    if not m:
        return t
    rest = m.group(2).replace('/', '\\')
    return f'{m.group(1).upper()}:/{rest}'


# ── Daten aus der Sicherung ─────────────────────────────────

@dataclass
class Bericht:
    hinweise: list[str] = field(default_factory=list)

    def __call__(self, meldung: str):
        self.hinweise.append(meldung)


def finde_datei(angabe: Path | None, muster: str, art: str, endung: str) -> Path:
    """Angegebene Datei oder – ohne Angabe – die neueste passende im aktuellen bzw. Programmordner."""
    ordner = list(dict.fromkeys([Path.cwd(), Path(__file__).resolve().parent]))
    kandidaten = sorted({p.resolve() for o in ordner for p in o.glob(muster)}, key=lambda p: p.stat().st_mtime, reverse=True)
    if angabe:
        # Windows blendet Endungen oft aus: „Datei.json“ heißt dann evtl. „Datei.json.json“ oder ohne Endung
        for versuch in (angabe, angabe.with_name(angabe.name + endung), angabe.with_suffix(endung)):
            if versuch.is_file():
                return versuch
        liste = ''.join(f'\n    {k.name}' for k in kandidaten[:10]) or '\n    (keine)'
        sys.exit(f'{art} „{angabe}“ nicht gefunden (gesucht in: {angabe.resolve().parent}).\n'
                 f'  Passende Dateien ({muster}):{liste}\n'
                 f'  Tipp: Dateinamen genau so angeben, wie er im Explorer steht, oder weglassen – '
                 f'dann wird automatisch die neueste passende Datei genommen.')
    if not kandidaten:
        sys.exit(f'Keine {art} gefunden ({muster}) in: ' + ', '.join(str(o) for o in ordner))
    print(f'{art}: {kandidaten[0].name}' + (f'  (neueste von {len(kandidaten)})' if len(kandidaten) > 1 else ''))
    return kandidaten[0]


def lade_sicherung(pfad: Path) -> dict:
    try:
        daten = json.loads(pfad.read_text(encoding='utf-8-sig'))
    except (UnicodeDecodeError, json.JSONDecodeError):
        sys.exit(f'„{pfad.name}“ ist keine lesbare JSON-Datei.')
    if daten.get('format') != 'wartungsprotokolle':
        sys.exit(f'„{pfad.name}“ ist keine Sicherung der Wartungs-App (Komplettsicherung aus „Daten“ verwenden).')
    return daten


def sektionen(plan, typ):
    return [s for s in plan or [] if s.get('typ') == typ]


def zaehlerstand(p, zeile_id):
    for sek in sektionen(p.get('plan'), 'tabelle'):
        werte = (p.get('werte') or {}).get(sek['id'], {}).get(zeile_id)
        if not werte:
            continue
        for spalte in ('nach', 'vor'):
            if zahl(werte.get(spalte)) is not None:
                return werte[spalte]
        for spalte in reversed(sek.get('spalten', [])):
            if zahl(werte.get(spalte['id'])) is not None:
                return werte[spalte['id']]
    return None


def arbeiten(p):
    erledigt, geplant = set(), False
    for sek in sektionen(p.get('plan'), 'aufgaben'):
        for el_id, w in (p.get('werte') or {}).get(sek['id'], {}).items():
            if w and w.get('erledigt'):
                erledigt.add(el_id)
            geplant |= bool(w and w.get('geplant'))
    return erledigt, geplant


def maengel(p) -> list[str]:
    liste = []
    werte = p.get('werte') or {}
    for sek in p.get('plan') or []:
        w = werte.get(sek['id']) or {}
        if sek.get('typ') == 'checkliste':
            for el in sek.get('elemente', []):
                v = w.get(el.get('id')) or {}
                if el.get('art') == 'ueberschrift' or v.get('s') != 'mangel':
                    continue
                eintrag = f"{sek.get('titel', '')}: {el.get('label', '')}"
                if (v.get('notiz') or '').strip():
                    eintrag += f" – {v['notiz'].strip()}"
                liste.append(eintrag + (' (behoben)' if v.get('behoben') else ''))
        elif sek.get('typ') == 'batterien':
            for g in sek.get('gruppen', []):
                for i, b in enumerate((w.get('gruppen') or {}).get(g['id'], []), 1):
                    if b.get('klemmen') == 'mangel':
                        liste.append(f"{sek.get('titel', '')}: {g.get('name', '')} {i} Anschlussklemmen"
                                     + (' (behoben)' if b.get('behoben') else ''))
    for m in p.get('maengel') or []:
        if (m.get('text') or '').strip():
            liste.append(m['text'].strip() + (' (behoben)' if m.get('behoben') else ''))
    return liste


def bemerkung(liste: list[str], allgemein) -> str:
    teile = []
    if liste:
        teile.append('Mängel: ' + ' / '.join(f'{i}) {text(m)}' for i, m in enumerate(liste, 1)))
    if text(allgemein):
        teile.append(text(allgemein))
    t = ' | '.join(teile)
    return t if len(t) <= MAX_ZELLE else t[:MAX_ZELLE - len(GEKUERZT)] + GEKUERZT


def techniker(name) -> list[str]:
    namen = [n.strip() for n in re.split(r'\s*(?:,|;|/|&|\bund\b)\s*', str(name or '')) if n.strip()]
    if len(namen) > 3:
        namen = namen[:2] + [', '.join(namen[2:])]
    return (namen + ['', '', ''])[:3]


# ── Bestehende Excel-Mappe (optional) ───────────────────────

@dataclass
class Bestand:
    aggregate: set[str] = field(default_factory=set)
    technik: set[str] = field(default_factory=set)
    technik_zeilen: list[dict] = field(default_factory=list)  # Zeilen von tblTechnik in Excel-Reihenfolge
    technik_start: int = 2  # Excel-Zeilennummer der ersten Datenzeile von tblTechnik
    kunden: dict[str, str] = field(default_factory=dict)   # normierter Name → Kunden-Nr
    kunden_original: dict[str, str] = field(default_factory=dict)
    mitarbeiter: set[str] = field(default_factory=set)
    termine: set[tuple[str, str]] = field(default_factory=set)  # (TT.MM.JJJJ, Kommission)
    reihenfolge: dict[str, list[str]] = field(default_factory=dict)  # Tabelle → Eingabespalten in Excel-Reihenfolge
    zuordnungen: dict[str, str] = field(default_factory=dict)  # Meldungen über abweichende Schreibweisen

    def kommission(self, roh) -> str:
        """Kommission so, wie sie in Excel steht.

        Altdaten haben Kommissionen teils als Zahl gespeichert, dabei gingen Endnullen verloren
        (23.050 → 23.05). Ist die Kommission in Excel unbekannt, wird sie über den Zahlenwert zugeordnet.
        """
        k = str(roh if roh is not None else '').strip()
        if not k or k in self.aggregate or not re.fullmatch(r'\d+(\.\d+)?', k):
            return k
        treffer = [e for e in self.aggregate if re.fullmatch(r'\d+(\.\d+)?', e) and float(e) == float(k)]
        if len(treffer) == 1:
            self.zuordnungen[k] = (f'Kommission „{k}“ der App als „{treffer[0]}“ aus Excel zugeordnet '
                                   f'(Endnullen fehlen in der App – dort korrigieren).')
            return treffer[0]
        return k

    def aehnliche_kommissionen(self, kommission: str) -> list[str]:
        import difflib
        k = kommission.casefold()
        teil = [e for e in self.aggregate if k and (e.casefold().startswith(k) or k.startswith(e.casefold()))]
        return list(dict.fromkeys(sorted(teil)[:3] + difflib.get_close_matches(kommission, self.aggregate, n=3, cutoff=0.8)))[:4]

    def kunde(self, name) -> tuple[str, str] | None:
        """(Kunden-Nr, Name in Excel); bei kleinen Abweichungen (z. B. „-rn“ aus verschluckten Zeilenumbrüchen) die ähnlichste."""
        import difflib
        n = norm_name(name)
        if not n:
            return None
        if n in self.kunden:
            return self.kunden[n], self.kunden_original[n]
        bewertet = sorted(((difflib.SequenceMatcher(None, n, k).ratio(), k) for k in self.kunden), reverse=True)
        if bewertet and bewertet[0][0] >= 0.9 and (len(bewertet) == 1 or bewertet[1][0] < bewertet[0][0] - 0.05):
            k = bewertet[0][1]
            self.zuordnungen[n] = (f'Kunde „{str(name).strip()}“ der App als „{self.kunden_original[k]}“ '
                                   f'[{self.kunden[k]}] zugeordnet – bitte prüfen.')
            return self.kunden[k], self.kunden_original[k]
        return None


# Eingabespalten laut PWA_Importformat.md, Abschnitt 6. Geschrieben wird in der Reihenfolge der Excel-Mappe.
EINGABEBLOECKE = {
    'tblTermine': ('Termine', [
        'Datum', 'Aggregat', 'Projekt-Nr', 'Status', 'Techniker 1', 'Techniker 2', 'Techniker 3',
        'Elektr. Wartung', 'Motorwartung', 'Ölwechsel', 'DGUV', 'Kühlmittel intern', 'Kühlmittel extern',
        'Luftfilter', 'Zahnriemen', 'Offene Arbeiten', 'Betriebsstunden', 'Startzahl', 'Bemerkung']),
    'tblAggregate': ('Aggregate', [
        'Kommission', 'Bezeichnung', 'Kunde', 'Standort', 'Straße', 'PLZ', 'Ort', 'Hersteller', 'Typ', 'Serien-Nr',
        'Leistung (kVA)', 'Motortyp', 'Motor-Nr', 'Kühlmittel intern Menge', 'Kühlmittel intern Sorte',
        'Kühlmittel extern Menge', 'Kühlmittel extern Sorte', 'Intervall Elektr. Wartung (Monate)',
        'Intervall Motorwartung (Monate)', 'Intervall Ölwechsel (Monate)', 'Intervall DGUV (Monate)',
        'Intervall Kühlmittel intern (Monate)', 'Intervall Kühlmittel extern (Monate)', 'Intervall Luftfilter (Monate)',
        'Wartungsvertrag', 'Vertrags-Nr', 'Wartungspauschale', 'Server-Link', 'Bemerkung']),
    'tblTechnik': ('Technische Daten', [
        'Aggregat', 'Generator Hersteller', 'Generator Typ', 'Generator Serien-Nr', 'Generator Regler Typ',
        'Motorregler Hersteller', 'Motorregler Typ', 'NSA Hersteller', 'NSA Typ', 'NSA Serien-Nr',
        'Ladegerät 1 Hersteller', 'Ladegerät 1 Typ', 'Ladegerät 2 Hersteller', 'Ladegerät 2 Typ',
        'Starterbatterie Hersteller', 'Starterbatterie Typ', 'Starterbatterie Anzahl', 'Starterbatterie Spannung (V)',
        'Starterbatterie Kapazität (Ah)', 'Starterbatterie wartungsfrei', 'Steuerbatterie Hersteller',
        'Steuerbatterie Typ', 'Steuerbatterie Anzahl', 'Steuerbatterie Spannung (V)', 'Steuerbatterie Kapazität (Ah)',
        'Steuerbatterie wartungsfrei']),
    'tblKunden': ('Kunden', ['Kunden-Nr', 'Name']),
    'tblMitarbeiter': ('Mitarbeiter', ['Name']),
}


def norm_name(t) -> str:
    return ' '.join(str(t or '').split()).casefold()


def norm_kopf(t) -> str:
    return ' '.join(str(t or '').replace('\xa0', ' ').split())


def pruefe_block(tabelle: str, blatt: str, kopf: list[str], soll: list[str]) -> list[str]:
    """Fehlertexte, wenn Eingabespalten fehlen oder fremde Spalten zwischen ihnen liegen (Reihenfolge ist egal)."""
    import difflib
    fehler = []
    ist = [norm_kopf(k) for k in kopf]
    fehlend = [sp for sp in soll if sp not in ist]
    if fehlend:
        def mit_vorschlag(sp):
            aehnlich = difflib.get_close_matches(sp, [k for k in ist if k not in soll], n=1, cutoff=0.6)
            return f'„{sp}“' + (f' (vorhanden ist „{aehnlich[0]}“)' if aehnlich else '')
        fehler.append(f'{len(fehlend)} von {len(soll)} Eingabespalten fehlen: ' + ', '.join(map(mit_vorschlag, fehlend)))
    doppelt = sorted({sp for sp in soll if ist.count(sp) > 1})
    for sp in doppelt:
        fehler.append(f'Spalte „{sp}“ kommt mehrfach vor')
    if fehlend or doppelt:
        return fehler
    pos = sorted(ist.index(sp) for sp in soll)
    im_block = [ist[i] for i in range(pos[0], pos[-1] + 1) if ist[i] not in soll]
    if im_block:
        fehler.append(
            f'Diese Spalten liegen zwischen den Eingabespalten und würden beim Einfügen überschrieben: '
            + ', '.join(f'„{sp}“' for sp in im_block)
            + f'.\n      Lösung: In Excel im Blatt „{blatt}“ jede dieser Spalten ganz nach rechts verschieben, '
              f'hinter „{ist[-1]}“\n      (Spaltenkopf anklicken, Umschalttaste gedrückt halten, am Rand der markierten '
              'Spalte anfassen und ziehen). Die Formeln funktionieren danach unverändert weiter.')
    return fehler


def lese_excel(pfad: Path, protokoll: Path) -> Bestand:
    try:
        from openpyxl import load_workbook
    except ImportError:
        sys.exit('Zum Lesen der Excel-Mappe wird openpyxl benötigt:  pip install openpyxl')
    try:
        wb = load_workbook(pfad, data_only=True)
    except Exception as e:  # beschädigt, falsches Format, in Excel gesperrt …
        sys.exit(f'Excel-Mappe „{pfad.name}“ kann nicht gelesen werden: {e}')

    tabellen, koepfe, fundorte, startzeilen = {}, {}, {}, {}
    for ws in wb.worksheets:
        for tbl in ws.tables.values():
            zeilen = list(ws[tbl.ref])
            kopf = [norm_kopf(z.value) for z in zeilen[0]]
            koepfe[tbl.displayName] = kopf
            fundorte[tbl.displayName] = f'Blatt „{ws.title}“, Bereich {tbl.ref}'
            startzeilen[tbl.displayName] = zeilen[0][0].row + 1
            tabellen[tbl.displayName] = [dict(zip(kopf, (z.value for z in zeile))) for zeile in zeilen[1:]]

    fehler = []
    for tabelle, (blatt, soll) in EINGABEBLOECKE.items():
        if tabelle not in koepfe:
            fehler.append(f'[{blatt}] Tabelle „{tabelle}“ nicht gefunden.')
            continue
        fehler += [f'[{blatt}] {f}' for f in pruefe_block(tabelle, blatt, koepfe[tabelle], soll)]

    diagnose = [f'Excel-Mappe: {pfad.resolve()}', '', 'Gefundene Blätter: ' + ', '.join(ws.title for ws in wb.worksheets), '']
    diagnose += ['Gefundene Tabellen und Spaltenköpfe:']
    for name, kopf in koepfe.items():
        diagnose.append(f'  {name}  ({fundorte[name]}, {len(tabellen[name])} Zeilen)')
        diagnose += [f'    {i:>2}. {k or "(leer)"}' for i, k in enumerate(kopf, 1)]
    if fehler:
        text_ = '\n'.join(['STRUKTURPRÜFUNG FEHLGESCHLAGEN – es wurden keine Importdateien erzeugt.', '',
                            'Probleme:', *[f'  - {f}' for f in fehler], '',
                            'Ist der Aufbau der Mappe so gewollt, diesen Bericht an den Entwickler schicken,',
                            'damit das Programm angepasst wird. Die Mappe wurde nicht verändert.', '',
                            *diagnose, ''])
        protokoll.write_text(text_, encoding='utf-8-sig', newline='\r\n')
        print(text_)
        sys.exit(f'Abgebrochen. Vollständiger Bericht: {protokoll}')
    print(f'Strukturprüfung bestanden ({len(EINGABEBLOECKE)} Tabellen).')

    def spalte(tabelle, name):
        return [z.get(name) for z in tabellen.get(tabelle, []) if z.get(name) not in (None, '')]

    b = Bestand()
    for tabelle, (_, soll) in EINGABEBLOECKE.items():
        b.reihenfolge[tabelle] = [k for k in koepfe[tabelle] if k in soll]
    b.aggregate = {str(k).strip() for k in spalte('tblAggregate', 'Kommission')}
    b.technik = {schluessel_aus_ref(a) for a in spalte('tblTechnik', 'Aggregat')}
    b.technik_zeilen = tabellen['tblTechnik']
    b.technik_start = startzeilen['tblTechnik']
    b.mitarbeiter = {str(n).strip() for n in spalte('tblMitarbeiter', 'Name')}
    for z in tabellen['tblKunden']:
        if z.get('Kunden-Nr') not in (None, '') and z.get('Name'):
            try:
                b.kunden[norm_name(z['Name'])] = str(round(float(z['Kunden-Nr'])))
                b.kunden_original[norm_name(z['Name'])] = str(z['Name']).strip()
            except (TypeError, ValueError):
                print(f'Achtung: Kunden-Nr „{z["Kunden-Nr"]}“ bei „{z["Name"]}“ ist keine Zahl – ignoriert.')
    for z in tabellen['tblTermine']:
        d = z.get('Datum')
        d = d.strftime('%d.%m.%Y') if isinstance(d, (datetime, date)) else str(d or '').strip()
        kommission = schluessel_aus_ref(z.get('Aggregat'))
        if d and kommission:
            b.termine.add((d, kommission))
    return b


# ── Zeilen je Blatt ─────────────────────────────────────────

def termin_zeile(p, anlage, bestand: Bestand, hinweis: Bericht) -> tuple[list[str], str] | None:
    s = (anlage or {}).get('stammdaten') or (p.get('anlage') or {}).get('stammdaten') or {}
    kommission = bestand.kommission(s.get('kommission'))
    if not kommission:
        hinweis(f"Protokoll vom {datum(p.get('datum'))} ohne Kommission – übersprungen.")
        return None
    erledigt, geplant = arbeiten(p)
    liste = maengel(p)
    meta = p.get('meta') or {}
    zeile = [
        datum(p.get('datum')),
        ref_aggregat(s.get('bezeichnung'), kommission),
        text(meta.get('auftrag')),
        'erledigt',
        *[text(t) for t in techniker(meta.get('techniker'))],
        *[x(l in erledigt) for l in LEISTUNGEN],
        x(bool(liste) or geplant),
        ganzzahl(zaehlerstand(p, 'betriebsstunden')),
        ganzzahl(zaehlerstand(p, 'startzaehler')),
        bemerkung(liste, p.get('bemerkung')),
    ]
    assert len(zeile) == 19
    return zeile, kommission


def aggregat_zeile(a, kommission, kunde: tuple[str, str]) -> list[str]:
    s, z, v = a.get('stammdaten') or {}, a.get('zusatz') or {}, a.get('vertrag') or {}
    zeile = [
        text(kommission), text(s.get('bezeichnung')), ref_kunde(kunde[1], kunde[0]),
        text(s.get('standort')), '', '', '',
        text(s.get('hersteller')), text(s.get('typ')), text(s.get('seriennr')),
        ganzzahl(z.get('Leistung_kVA')), text(z.get('Motor_Typ')), text(z.get('Motor_NR')),
        ganzzahl(z.get('Menge_Kuehlmittel_intern')), text(z.get('Typ_Kuehlmittel_intern')),
        ganzzahl(z.get('Menge_Kuehlmittel_extern')), text(z.get('Typ_Kuehlmittel_extern')),
        *[ganzzahl(z.get(i)) for i in INTERVALLE],
        ja_nein(bool(v.get('aktiv'))), text(v.get('nummer')), geld(z.get('Wartungspauschale')),
        text(server_link(z.get('Server_Link'))), text(a.get('notizen')),
    ]
    assert len(zeile) == 29
    return zeile


def excel_wert(v) -> str:
    """Zellwert aus der Mappe im Format der Importdateien (zum Zurückschreiben unveränderter Werte)."""
    if v is None:
        return ''
    if isinstance(v, bool):
        return 'Ja' if v else 'Nein'
    if isinstance(v, (datetime, date)):
        return v.strftime('%d.%m.%Y')
    if isinstance(v, (int, float)):
        return dezimal(v)
    return text(v)


# Felder (Index in technik_zeile), die ohne Hersteller/Typ/Kapazität nur Vorgabewerte der Vorlage sind
BATTERIE_VORGABE = {16: (14, 15, 18), 17: (14, 15, 18), 19: (14, 15, 18),
                    22: (20, 21, 24), 23: (20, 21, 24), 25: (20, 21, 24)}


def technik_ersetzen(bestand: 'Bestand', app_zeilen: dict[str, list[str]]):
    """Kompletter Eingabeblock von „Technische Daten“ mit Werten aus der App.

    Excel kann bestehende Zeilen beim Einfügen nicht aktualisieren. Deshalb werden alle vorhandenen
    Zeilen in unveränderter Reihenfolge ausgegeben und über den alten Block eingefügt. Leere Werte der App
    überschreiben nichts. Liefert (Zeilen in Excel-Spaltenfolge, Änderungen) oder (None, []) ohne Änderung.
    """
    soll = EINGABEBLOECKE['tblTechnik'][1]
    folge = bestand.reihenfolge['tblTechnik']
    zeilen, aenderungen = [], []
    for nr, excel in enumerate(bestand.technik_zeilen, bestand.technik_start):
        alt = [excel_wert(excel.get(sp)) for sp in soll]
        neu = list(alt)
        kommission = schluessel_aus_ref(excel.get('Aggregat'))
        app = app_zeilen.get(kommission)
        if app:
            for i in range(1, len(soll)):
                if not app[i] or app[i] == alt[i]:
                    continue
                if i in BATTERIE_VORGABE and not any(app[j] for j in BATTERIE_VORGABE[i]):
                    continue
                neu[i] = app[i]
                aenderungen.append([str(nr), kommission, soll[i], alt[i], app[i]])
        zeilen.append([neu[soll.index(sp)] for sp in folge])
    return (zeilen if aenderungen else None), aenderungen


def technik_zeile(a, kommission) -> list[str] | None:
    s, z = a.get('stammdaten') or {}, a.get('zusatz') or {}
    bat = next(iter(sektionen(a.get('pruefplan'), 'batterien')), {})
    gruppen = {g.get('id'): g for g in bat.get('gruppen', [])}
    lader = [l.get('name', '') if l.get('aktiv', True) else '' for l in bat.get('lader', [])] + ['', '']
    lader = ['' if re.fullmatch(r'Ladegerät \d', n or '') else n for n in lader]

    def batterie(gid):
        g = gruppen.get(gid)
        if not g or not zahl(g.get('anzahl')):
            return [''] * 6
        return [text(g.get('hersteller')), text(g.get('typ')), ganzzahl(g.get('anzahl')),
                ganzzahl(g.get('spannung')), ganzzahl(g.get('kapazitaet')), ja_nein(bool(g.get('wartungsfrei')))]

    zeile = [
        ref_aggregat(s.get('bezeichnung'), kommission),
        *[text(z.get(k)) for k in ('Generator_Hersteller', 'Generator_Typ', 'Generator_Serien_NR', 'Generator_Regler_Typ',
                                   'Motorregler_Hersteller', 'Motorregler_Typ', 'NSA_Hersteller', 'NSA_Typ', 'NSA_Serien_NR')],
        text(lader[0]), '', text(lader[1]), '',
        *batterie('starterbatterie'), *batterie('steuerbatterie'),
    ]
    assert len(zeile) == 26
    # Ohne Gerätedaten (nur Standardwerte der Vorlage wie 2 × 12 V) keine Zeile anlegen
    hat_daten = any(zeile[1:14]) or any(zeile[i] for i in (14, 15, 18, 20, 21, 24))
    return zeile if hat_daten else None


# ── Ausgabe ─────────────────────────────────────────────────

def schreibe(ordner: Path, name: str, zeilen: list[list[str]]) -> Path | None:
    if not zeilen:
        return None
    pfad = ordner / name
    with pfad.open('w', encoding='utf-8-sig', newline='') as f:
        f.write(''.join('\t'.join(z) + '\r\n' for z in zeilen))
    return pfad


def schreibe_xlsx(pfad: Path, blaetter: list[tuple[str, list[str], list[list[str]]]]):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)
    for name, kopf, zeilen in blaetter:
        ws = wb.create_sheet(name)
        ws.append(kopf)
        for zeile in zeilen:
            ws.append([None] * len(kopf))
            for spalte, wert in enumerate(zeile, 1):
                zelle = ws.cell(row=ws.max_row, column=spalte)
                # Als Text ablegen: führende Nullen bleiben, „=…“ wird keine Formel
                zelle.value = wert[1:] if wert.startswith("'") else wert
                zelle.data_type = 's'
                zelle.number_format = '@'
        for zelle in ws[1]:
            zelle.font = Font(bold=True)
            zelle.fill = PatternFill('solid', fgColor='DDE3EA')
            zelle.alignment = Alignment(wrap_text=True, vertical='top')
        for i, k in enumerate(kopf, 1):
            breite = max([len(k)] + [len(str(z[i - 1])) for z in zeilen])
            ws.column_dimensions[get_column_letter(i)].width = min(max(breite + 2, 10), 60)
        ws.freeze_panes = 'B2'
        ws.auto_filter.ref = ws.dimensions
    wb.save(pfad)


ANLEITUNG = """ÜBERNAHME IN „Wartungsdaten.xlsx“
====================================

Für jede Datei, in dieser Reihenfolge:
{dateien}
So wird eine Datei eingefügt:
  1. Die Datei mit dem Editor öffnen (Doppelklick).
  2. Strg+A (alles markieren), dann Strg+C (kopieren).
  3. In Excel das genannte Blatt öffnen.
  4. In die erste leere Zeile direkt unter der Tabelle klicken, und zwar in die oben
     genannte Spalte.
     Die Excel-Mappe danach nicht mehr umsortieren, sonst das Programm neu laufen lassen.
  5. Strg+V (einfügen). Die Tabelle wächst automatisch, Formelspalten füllen sich selbst.
  6. Kurz prüfen: stehen die Werte in den richtigen Spalten? Wenn nicht: Strg+Z und Schritt 4 prüfen.

Jede Datei nur EIN Mal einfügen – Excel erkennt keine doppelten Zeilen.
{ersetzen}"""

ERSETZEN = """

TECHNISCHE DATEN ERSETZEN (vorhandene Zeilen aktualisieren)
===========================================================

„{datei}“ enthält ALLE {zeilen} vorhandenen Zeilen des Blatts „Technische Daten“ in der
bisherigen Reihenfolge – mit {anzahl} geänderten Werten bei {aggregate} Aggregaten aus der App
(Liste zum Prüfen: „Aenderungen_Technische_Daten_….xlsx“). Diese Datei wird NICHT unter die Tabelle,
sondern ÜBER die vorhandenen Zeilen eingefügt:

  1. Vorher die Excel-Mappe sichern (Kopie anlegen).
  2. Im Blatt „Technische Daten“ alle Filter aufheben und NICHT sortieren – die Zeilen müssen
     genau so stehen wie beim Programmlauf. Zeile {erste_zeile}: „{erster}“ … Zeile {letzte_zeile}: „{letzter}“.
  3. Datei im Editor öffnen, Strg+A, Strg+C.
  4. In Excel in Zeile {erste_zeile} (erste Datenzeile) in die Spalte „{spalte}“ klicken und Strg+V.
  5. Prüfen: In Spalte „{spalte}“ steht in jeder Zeile noch dasselbe Aggregat wie vorher.
     Wenn nicht: sofort Strg+Z.

Erst danach die übrigen Dateien einfügen (Technische Daten neuer Aggregate kommen dann unten dazu).
"""


def main():
    ap = argparse.ArgumentParser(description='Sicherung der Wartungs-App → Importdateien für Wartungsdaten.xlsx')
    ap.add_argument('sicherung', type=Path, nargs='?',
                    help='Komplettsicherung der App; ohne Angabe die neueste Wartung_Backup*.json im Ordner')
    ap.add_argument('--excel', type=Path, help='Arbeitsmappe (wird vorher geprüft); ohne Angabe Wartungsdaten*.xlsx im Ordner')
    ap.add_argument('--ziel', type=Path, help='Ausgabeordner (Standard: neben der Sicherung)')
    args = ap.parse_args()

    sicherung = finde_datei(args.sicherung, 'Wartung_Backup*.json', 'Sicherung', '.json')
    excel = finde_datei(args.excel, 'Wartungsdaten*.xlsx', 'Excel-Mappe', '.xlsx')
    daten = lade_sicherung(sicherung)
    heute = date.today().isoformat()
    ordner = args.ziel or sicherung.with_name(f'Excel-Import_{heute}')
    ordner.mkdir(parents=True, exist_ok=True)
    bestand = lese_excel(excel, ordner / 'Strukturpruefung.txt')
    hinweis = Bericht()

    anlagen = {a['id']: a for a in daten.get('anlagen', [])}
    protokolle = sorted(daten.get('protokolle', []), key=lambda p: (p.get('datum') or '', p.get('id')))

    # Aggregate und Technische Daten. Nicht importierbare Anlagen kommen mit allen Daten in eine eigene Mappe.
    aggregate, technik, neue_kommissionen = [], [], set()
    offen = {'tblAggregate': [], 'tblTechnik': [], 'tblTermine': []}
    offen_kommissionen: dict[str, str] = {}
    technik_app: dict[str, list[str]] = {}
    for a in sorted(anlagen.values(), key=lambda a: str(a.get('stammdaten', {}).get('kommission') or '')):
        s = a.get('stammdaten') or {}
        kommission = bestand.kommission(s.get('kommission'))
        if not kommission:
            continue
        if kommission not in bestand.aggregate:
            kunde = bestand.kunde(s.get('kunde'))
            if not kunde:
                aehnlich = bestand.aehnliche_kommissionen(kommission)
                grund = (f'Kunde „{s["kunde"].strip()}“ fehlt im Blatt „Kunden“ (dort anlegen, Name wie in der App)'
                         if str(s.get('kunde') or '').strip() else 'in der App ist kein Kunde eingetragen')
                grund = (f'Neu für Excel, aber {grund}.'
                         + (f' Ähnliche Kommissionen in Excel: {", ".join(aehnlich)}.' if aehnlich else ''))
                offen_kommissionen[kommission] = grund
                zeile = aggregat_zeile(a, kommission, ('?', ''))
                zeile[2] = text(s.get('kunde'))
                offen['tblAggregate'].append([grund, *zeile])
                if zeile_t := technik_zeile(a, kommission):
                    offen['tblTechnik'].append([grund, *zeile_t])
                continue
            aggregate.append(aggregat_zeile(a, kommission, kunde))
            neue_kommissionen.add(kommission)
        if zeile := technik_zeile(a, kommission):
            if kommission in bestand.technik:
                technik_app[kommission] = zeile
            else:
                technik.append(zeile)

    # Termine
    termine, entwuerfe, doppelt, gesehen = [], 0, 0, set()
    namen = []
    for p in protokolle:
        if p.get('status') != 'abgeschlossen':
            entwuerfe += 1
            continue
        ergebnis = termin_zeile(p, anlagen.get(p.get('anlageId')), bestand, hinweis)
        if not ergebnis:
            continue
        zeile, kommission = ergebnis
        if (zeile[0], kommission) in bestand.termine:
            doppelt += 1
            continue
        if (zeile[0], kommission) in gesehen:
            hinweis(f'Mehrere Protokolle am {zeile[0]} für {kommission} – nur das erste exportiert.')
            continue
        gesehen.add((zeile[0], kommission))
        if kommission not in bestand.aggregate and kommission not in neue_kommissionen:
            offen['tblTermine'].append([offen_kommissionen.get(kommission, 'Aggregat fehlt in Excel.'), *zeile])
            continue
        termine.append(zeile)
        namen += [n for n in zeile[4:7] if n and ',' not in n]

    mitarbeiter = [[n] for n in dict.fromkeys(namen) if n not in bestand.mitarbeiter]

    dateien = []
    for nr, (tabelle, zeilen) in enumerate([('tblMitarbeiter', mitarbeiter), ('tblAggregate', aggregate),
                                            ('tblTechnik', technik), ('tblTermine', termine)], 1):
        blatt, soll = EINGABEBLOECKE[tabelle]
        excel_folge = bestand.reihenfolge[tabelle]
        zeilen = [[z[soll.index(sp)] for sp in excel_folge] for z in zeilen]
        pfad = schreibe(ordner, f'Import_{nr}_{blatt.replace(" ", "_")}_{heute}.txt', zeilen)
        if pfad:
            dateien.append((blatt, pfad, len(zeilen), excel_folge[0]))

    ersatz, aenderungen = technik_ersetzen(bestand, technik_app)
    if ersatz:
        pfad = schreibe(ordner, f'Ersetzen_Technische_Daten_{heute}.txt', ersatz)
        schreibe_xlsx(ordner / f'Aenderungen_Technische_Daten_{heute}.xlsx',
                      [('Änderungen', ['Excel-Zeile', 'Kommission', 'Spalte', 'Bisher in Excel', 'Neu aus der App'], aenderungen)])
        erste = bestand.reihenfolge['tblTechnik'][0]
        ersetzen_text = ERSETZEN.format(datei=pfad.name, zeilen=len(ersatz), spalte=erste,
                                        erster=excel_wert(bestand.technik_zeilen[0].get('Aggregat')),
                                        letzter=excel_wert(bestand.technik_zeilen[-1].get('Aggregat')),
                                        erste_zeile=bestand.technik_start,
                                        letzte_zeile=bestand.technik_start + len(ersatz) - 1, anzahl=len(aenderungen),
                                        aggregate=len({a[1] for a in aenderungen}))
        hinweis(f'Technische Daten: {len(aenderungen)} Werte bei {len({a[1] for a in aenderungen})} vorhandenen '
                f'Aggregaten weichen von der App ab – siehe „Aenderungen_Technische_Daten_{heute}.xlsx“ und Abschnitt '
                f'„Technische Daten ersetzen“ in der Anleitung.')
    else:
        ersetzen_text = ''

    if any(offen.values()):
        pfad = ordner / f'Nicht_importiert_{heute}.xlsx'
        blaetter = []
        for t, zeilen in offen.items():
            if zeilen:
                blatt, soll = EINGABEBLOECKE[t]
                folge = bestand.reihenfolge[t]
                blaetter.append((blatt, ['Grund', *folge], [[z[0], *(z[1 + soll.index(sp)] for sp in folge)] for z in zeilen]))
        schreibe_xlsx(pfad, blaetter)
        hinweis(f'{len(offen_kommissionen)} Anlage(n) mit {len(offen["tblTermine"])} Termin(en) passen nicht automatisch '
                f'nach Excel – alle Daten stehen in „{pfad.name}“ zum Eintragen von Hand (Spalte „Grund“ erklärt, warum).')

    (ordner / 'Anleitung.txt').write_text(ANLEITUNG.format(
        ersetzen=ersetzen_text,
        dateien=''.join(f'  {i}. {p.name}\n     → Blatt „{b}“, Spalte „{erste}“ ({n} Zeile{"" if n == 1 else "n"})\n'
                        for i, (b, p, n, erste) in enumerate(dateien, 1)),
    ), encoding='utf-8-sig', newline='\r\n')

    for meldung in bestand.zuordnungen.values():
        hinweis(meldung)
    hersteller = sorted({z[7] for z in aggregate if z[7]})
    if hersteller:
        hinweis(f'Hersteller müssen im Blatt „Listen“ unter „Aggregathersteller“ stehen (sonst dort ergänzen): '
                f'{", ".join(hersteller)}')
    if entwuerfe:
        hinweis(f'{entwuerfe} nicht abgeschlossene Protokolle (Entwürfe) übersprungen.')
    if doppelt:
        hinweis(f'{doppelt} Termine sind in Excel schon vorhanden und wurden übersprungen.')
    if hinweis.hinweise:
        (ordner / 'Hinweise.txt').write_text(''.join(f'- {h}\n' for h in hinweis.hinweise), encoding='utf-8-sig', newline='\r\n')

    print(f'Ausgabe: {ordner}')
    for blatt, pfad, n, _ in dateien:
        print(f'  {pfad.name}: {n} Zeile{"" if n == 1 else "n"}')
    for h in hinweis.hinweise:
        print(f'  Hinweis: {h}')


if __name__ == '__main__':
    main()
